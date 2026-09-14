from __future__ import annotations
from pathlib import Path
from datetime import datetime
import re
from openpyxl import load_workbook

def _is_wrong(v) -> bool:
    if v is None or v == "":
        return False
    if isinstance(v, (int, float)):
        return float(v) != 0
    s = str(v).strip().lower()
    return s not in {"", "0", "정답", "o", "○", "false", "none"}

def _date_from_name(name: str) -> str | None:
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(\d{2})(?!\d)", name)
    if not m:
        return None
    yy, mm, dd = map(int, m.groups())
    try:
        return datetime(2000 + yy, mm, dd).strftime("%Y-%m-%d")
    except ValueError:
        return None

def read_workbook(path: str) -> dict:
    wb = load_workbook(path, data_only=False)
    if "테스트" not in wb.sheetnames:
        raise ValueError("엑셀에서 '테스트' 시트를 찾지 못했습니다.")
    ws = wb["테스트"]

    header_row = name_col = None
    for r in range(1, min(ws.max_row, 30) + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value or "").strip() == "이름":
                header_row, name_col = r, c
                break
        if header_row:
            break
    if not header_row:
        raise ValueError("'테스트' 시트에서 '이름' 열을 찾지 못했습니다.")

    problem_cols = {}
    for c in range(name_col + 1, ws.max_column + 1):
        try:
            q = int(ws.cell(header_row, c).value)
            if 1 <= q <= 300:
                problem_cols[q] = c
        except Exception:
            pass
    if not problem_cols:
        raise ValueError("문제번호 열(1,2,3...)을 찾지 못했습니다.")

    question_count = 0
    if "테스트정보" in wb.sheetnames:
        info = wb["테스트정보"]
        for r in range(1, info.max_row + 1):
            try:
                q = int(info.cell(r, 2).value)
            except Exception:
                continue
            typ, area = info.cell(r, 3).value, info.cell(r, 4).value
            if typ not in (None, "") or area not in (None, ""):
                question_count = max(question_count, q)

    students = []
    max_wrong = 0
    for r in range(header_row + 1, ws.max_row + 1):
        name = str(ws.cell(r, name_col).value or "").strip()
        if not name:
            continue
        wrongs = []
        for q, c in problem_cols.items():
            if _is_wrong(ws.cell(r, c).value):
                wrongs.append(q)
                max_wrong = max(max_wrong, q)
        students.append({"name": name, "wrongs": wrongs})

    if question_count == 0:
        question_count = max_wrong or max(problem_cols)

    stem = Path(path).stem
    return {
        "students": students,
        "question_count": question_count,
        "test_name": stem,
        "test_date": _date_from_name(stem) or datetime.now().strftime("%Y-%m-%d"),
    }
