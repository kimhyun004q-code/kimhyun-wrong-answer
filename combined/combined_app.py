from __future__ import annotations

import contextlib
import os
import re
import runpy
import shutil
import sys
import tempfile
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from tkinterdnd2 import TkinterDnD, DND_FILES
from openpyxl import load_workbook

from excel_reader import read_workbook
from hwpx_direct import HwpxExam

APP_TITLE = "김현수학 통합 성적표 + 개인별 오답 생성기"
VERSION = "v1.1 안정판"
DATE_RE = re.compile(r"(?<!\d)(\d{6})(?!\d)")
ROUND_RE = re.compile(r"(\d+)\s*(회차|회|차시)")


def bundle_dir() -> Path:
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))


def safe_name(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", str(text)).strip().rstrip(".") or "결과"


def normalize_round(text: str) -> str:
    m = ROUND_RE.search(text or "")
    return f"{int(m.group(1))}회차" if m else ""


def round_number(text: str) -> int:
    m = ROUND_RE.search(text or "")
    return int(m.group(1)) if m else 9999


def date_key_from_name(path: str) -> str:
    m = DATE_RE.search(Path(path).stem)
    return m.group(1) if m else ""


def excel_date_key(path: str) -> str:
    key = date_key_from_name(path)
    if key:
        return key
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "명단" in wb.sheetnames:
            ws = wb["명단"]
            for r in range(1, min(ws.max_row, 25) + 1):
                if str(ws.cell(r, 2).value or "").strip() == "날짜":
                    v = ws.cell(r, 3).value
                    if isinstance(v, datetime):
                        return v.strftime("%y%m%d")
                    s = str(v or "")
                    nums = re.findall(r"\d+", s)
                    if len(nums) >= 3:
                        y, m, d = map(int, nums[:3])
                        if y < 100:
                            y += 2000
                        return f"{y%100:02d}{m:02d}{d:02d}"
        wb.close()
    except Exception:
        pass
    return ""


def display_date(key: str) -> str:
    try:
        return datetime.strptime(key, "%y%m%d").strftime("%Y-%m-%d")
    except Exception:
        return key or "-"


def workbook_meta(path: str) -> dict:
    out = {"학원": "", "반": "", "날짜": ""}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "명단" in wb.sheetnames:
            ws = wb["명단"]
            for r in range(1, min(ws.max_row, 25) + 1):
                k = str(ws.cell(r, 2).value or "").strip()
                if k in out:
                    out[k] = str(ws.cell(r, 3).value or "").strip()
        wb.close()
    except Exception:
        pass
    return out


@dataclass
class Session:
    uid: str
    selected: bool
    date_key: str
    round_name: str
    xlsx: str
    hwpx: str


class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1120x820")
        self.minsize(960, 700)
        self.sessions: list[Session] = []
        self.files: list[str] = []
        self.output_dir = tk.StringVar(value=str(Path.home() / "Desktop" / "김현수학_통합결과"))
        self.status = tk.StringVar(value="엑셀과 HWPX 파일을 창에 끌어다 놓으세요.")
        self.last_result: Path | None = None
        self.running = False
        self._build()

    def _build(self):
        outer = ttk.Frame(self, padding=14)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text=APP_TITLE, font=("맑은 고딕", 18, "bold")).pack(anchor="w")
        ttk.Label(outer, text=f"{VERSION}  ·  선택한 1/2/3/여러 회차를 한 번에 처리", foreground="#667085").pack(anchor="w", pady=(2, 10))

        drop = tk.Label(
            outer,
            text="여기에 회차별 성적 XLSX + 시험지 HWPX를 모두 드래그하세요\n여러 파일을 한꺼번에 놓아도 됩니다.",
            relief="groove", bd=2, height=4, font=("맑은 고딕", 12, "bold"),
        )
        drop.pack(fill="x", pady=(0, 10))
        drop.drop_target_register(DND_FILES)
        drop.dnd_bind("<<Drop>>", self._on_drop)

        bar = ttk.Frame(outer)
        bar.pack(fill="x", pady=(0, 8))
        ttk.Button(bar, text="파일 추가", command=self._pick_files).pack(side="left")
        ttk.Button(bar, text="목록 비우기", command=self._clear).pack(side="left", padx=6)
        ttk.Button(bar, text="전체 선택", command=lambda: self._set_all(True)).pack(side="left", padx=(14, 4))
        ttk.Button(bar, text="전체 해제", command=lambda: self._set_all(False)).pack(side="left")
        ttk.Label(bar, text="   회차 행을 더블클릭하면 선택/해제").pack(side="left")

        cols = ("sel", "round", "date", "xlsx", "hwpx", "state")
        self.tree = ttk.Treeview(outer, columns=cols, show="headings", height=11)
        specs = [
            ("sel", "선택", 60), ("round", "회차", 90), ("date", "날짜", 105),
            ("xlsx", "성적 XLSX", 310), ("hwpx", "시험지 HWPX", 310), ("state", "상태", 120),
        ]
        for c, t, w in specs:
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor="center" if c in {"sel", "round", "date", "state"} else "w")
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<Double-1>", self._toggle_row)

        out = ttk.LabelFrame(outer, text="결과", padding=8)
        out.pack(fill="x", pady=(10, 8))
        ttk.Label(out, text="저장 폴더").grid(row=0, column=0, sticky="w")
        ttk.Entry(out, textvariable=self.output_dir).grid(row=0, column=1, sticky="ew", padx=8)
        ttk.Button(out, text="선택", command=self._pick_out).grid(row=0, column=2)
        out.columnconfigure(1, weight=1)
        ttk.Label(
            out,
            text="출력: ① 선택 회차 누적 성적표  ② 회차별 개인 오답 HWPX  ③ 학생별 폴더에 성적표 이미지+오답노트 묶음",
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(6, 0))

        actions = ttk.Frame(outer)
        actions.pack(fill="x", pady=(2, 6))
        self.btn = ttk.Button(actions, text="선택 회차 통합 생성", command=self._start)
        self.btn.pack(side="left")
        ttk.Button(actions, text="결과 폴더 열기", command=self._open_result).pack(side="left", padx=8)
        self.pb = ttk.Progressbar(actions, mode="determinate", maximum=100)
        self.pb.pack(side="right", fill="x", expand=True, padx=(20, 0))

        ttk.Label(outer, textvariable=self.status).pack(anchor="w")
        self.log = tk.Text(outer, height=8, wrap="word")
        self.log.pack(fill="x", pady=(4, 0))

    def _log(self, text: str):
        def apply():
            self.status.set(text)
            self.log.insert("end", time.strftime("%H:%M:%S ") + text + "\n")
            self.log.see("end")
        self.after(0, apply)

    def _progress(self, value: float):
        self.after(0, lambda: self.pb.configure(value=max(0, min(100, value))))

    def _on_drop(self, event):
        self._add_files(list(self.tk.splitlist(event.data)))

    def _pick_files(self):
        files = filedialog.askopenfilenames(filetypes=[("성적/시험지", "*.xlsx *.hwpx"), ("Excel", "*.xlsx"), ("HWPX", "*.hwpx")])
        if files:
            self._add_files(list(files))

    def _add_files(self, paths: list[str]):
        changed = False
        for p in paths:
            p = str(Path(p))
            if Path(p).suffix.lower() in {".xlsx", ".hwpx"} and p not in self.files:
                self.files.append(p); changed = True
        if changed:
            self._rebuild_sessions()

    def _clear(self):
        self.files.clear(); self.sessions.clear(); self._refresh_tree()
        self._log("목록을 비웠습니다.")

    def _pair_score(self, xlsx: str, hwpx: str) -> int:
        xd, hd = excel_date_key(xlsx), date_key_from_name(hwpx)
        xr, hr = normalize_round(Path(xlsx).stem), normalize_round(Path(hwpx).stem)
        score = 0
        if xd and hd and xd == hd: score += 100
        if xr and hr and xr == hr: score += 40
        if Path(xlsx).stem.lower() == Path(hwpx).stem.lower(): score += 20
        return score

    def _rebuild_sessions(self):
        old = {(s.date_key, s.round_name, Path(s.xlsx).name): s.selected for s in self.sessions}
        excels = sorted([p for p in self.files if Path(p).suffix.lower() == ".xlsx"], key=lambda p: (excel_date_key(p), round_number(Path(p).stem), Path(p).name))
        hws = [p for p in self.files if Path(p).suffix.lower() == ".hwpx"]
        unused = set(hws)
        sessions: list[Session] = []
        for idx, x in enumerate(excels, 1):
            key = excel_date_key(x)
            label = normalize_round(Path(x).stem)
            candidates = [(self._pair_score(x, h), h) for h in unused]
            candidates.sort(reverse=True, key=lambda z: z[0])
            h = candidates[0][1] if candidates and candidates[0][0] > 0 else ""
            if not h and len(unused) == 1:
                h = next(iter(unused))
            if h:
                unused.discard(h)
                if not label:
                    label = normalize_round(Path(h).stem)
            if not label:
                label = f"{idx}회차"
            uid = f"{key}|{label}|{Path(x).name}"
            sessions.append(Session(uid, old.get((key, label, Path(x).name), True), key, label, x, h))
        self.sessions = sessions
        self._refresh_tree()
        paired = sum(bool(s.hwpx) for s in sessions)
        self._log(f"회차 {len(sessions)}개 인식 · HWPX 짝맞춤 {paired}/{len(sessions)}")

    def _refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        for i, s in enumerate(self.sessions):
            state = "준비완료" if s.xlsx and s.hwpx else "HWPX 필요"
            self.tree.insert("", "end", iid=str(i), values=("☑" if s.selected else "☐", s.round_name, display_date(s.date_key), Path(s.xlsx).name if s.xlsx else "-", Path(s.hwpx).name if s.hwpx else "-", state))

    def _toggle_row(self, event=None):
        item = self.tree.identify_row(event.y) if event else ""
        if item:
            i = int(item); self.sessions[i].selected = not self.sessions[i].selected; self._refresh_tree()

    def _set_all(self, value: bool):
        for s in self.sessions: s.selected = value
        self._refresh_tree()

    def _pick_out(self):
        p = filedialog.askdirectory()
        if p: self.output_dir.set(p)

    def _selected(self) -> list[Session]:
        return [s for s in self.sessions if s.selected]

    def _start(self):
        if self.running: return
        selected = self._selected()
        if not selected:
            messagebox.showwarning("회차 선택", "처리할 회차를 하나 이상 선택해 주세요."); return
        missing = [s.round_name for s in selected if not s.hwpx]
        if missing:
            messagebox.showwarning("HWPX 확인", "시험지 HWPX가 연결되지 않은 회차가 있습니다:\n" + ", ".join(missing)); return
        self.running = True; self.btn.configure(state="disabled"); self.pb.configure(value=0)
        threading.Thread(target=self._run, args=(selected,), daemon=True).start()

    def _run_score_engine(self, selected: list[Session], score_dir: Path):
        # 안정판: EXE가 자기 자신을 자식 프로세스로 다시 실행하지 않는다.
        # 일부 PC/백신에서 one-file EXE의 self-spawn이 차단되어 빈 결과가 생기는 문제를 줄인다.
        class _ScoreLogStream:
            def __init__(self, emit):
                self.emit = emit
                self.buf = ""

            def write(self, text):
                if not text:
                    return 0
                self.buf += str(text).replace("\r", "")
                while "\n" in self.buf:
                    line, self.buf = self.buf.split("\n", 1)
                    line = line.strip()
                    if line:
                        self.emit(line)
                return len(text)

            def flush(self):
                line = self.buf.strip()
                self.buf = ""
                if line:
                    self.emit(line)

            def reconfigure(self, **kwargs):
                return None

        with tempfile.TemporaryDirectory(prefix="kimhyun_score_input_") as td:
            inp = Path(td)
            used = set()
            for i, s in enumerate(selected, 1):
                key = s.date_key or datetime.now().strftime("%y%m%d")
                name = Path(s.xlsx).name
                if not re.match(r"^\d{6}(?=[\s_.\-]|$)", Path(name).stem + "."):
                    name = f"{key}_{s.round_name}_{name}"
                if name in used:
                    name = f"{key}_{i:02d}_{s.round_name}_{name}"
                used.add(name)
                shutil.copy2(s.xlsx, inp / name)
            self._log("누적 성적표 생성 중…")
            stream = _ScoreLogStream(lambda line: self._log("[성적표] " + line))
            try:
                with contextlib.redirect_stdout(stream), contextlib.redirect_stderr(stream):
                    run_score_engine([str(inp), str(score_dir), str(len(selected))])
            except SystemExit as e:
                code = e.code
                if code not in (None, 0):
                    raise RuntimeError(f"성적표 생성 단계가 중단되었습니다. ({code})") from None
            finally:
                stream.flush()

    def _create_wrong_notes(self, selected: list[Session], wrong_root: Path, student_root: Path):
        total_sessions = len(selected)
        for si, s in enumerate(selected, 1):
            self._log(f"{s.round_name} 오답노트 분석 중…")
            data = read_workbook(s.xlsx)
            meta = workbook_meta(s.xlsx)
            test_date = data.get("test_date") or display_date(s.date_key)
            class_name = meta.get("반", "")
            round_dir = wrong_root / safe_name(s.round_name)
            round_dir.mkdir(parents=True, exist_ok=True)
            exam = HwpxExam(s.hwpx, data["question_count"])
            targets = [x for x in data["students"] if x.get("wrongs")]
            for j, st in enumerate(targets, 1):
                nm = str(st["name"]).strip()
                out = round_dir / f"{safe_name(nm)}_{safe_name(s.round_name)}_오답.hwpx"
                exam.generate(str(out), nm, test_date, st["wrongs"], class_name=class_name, round_name=s.round_name)
                person = student_root / safe_name(nm)
                person.mkdir(parents=True, exist_ok=True)
                shutil.copy2(out, person / out.name)
                frac = ((si - 1) + (j / max(1, len(targets)))) / max(1, total_sessions)
                self._progress(45 + frac * 45)
            self._log(f"{s.round_name} 개인별 오답 {len(targets)}명 완료")

    def _copy_score_images_to_students(self, score_dir: Path, student_root: Path):
        image_dirs = [p for p in score_dir.iterdir() if p.is_dir() and "이미지" in p.name]
        if not image_dirs: return
        imgdir = sorted(image_dirs, key=lambda p: p.stat().st_mtime)[-1]
        for img in imgdir.glob("*.png"):
            stem = re.sub(r"^\d+\s+", "", img.stem)
            stem = re.sub(r"\s+\(\d+\)$", "", stem).strip()
            if not stem: continue
            # 성적표 파일명 끝의 x는 '평균 대비 -20점 이하' 표시일 뿐 학생 이름에는 포함하지 않는다.
            student_name = stem[:-1].strip() if stem.endswith("x") else stem
            if not student_name: continue
            person = student_root / safe_name(student_name)
            person.mkdir(parents=True, exist_ok=True)
            shutil.copy2(img, person / "누적성적표.png")

    def _run(self, selected: list[Session]):
        try:
            root = Path(self.output_dir.get()).expanduser()
            root.mkdir(parents=True, exist_ok=True)
            tag = "_".join(s.round_name for s in selected)
            result = root / f"통합결과_{safe_name(tag)}_{datetime.now().strftime('%y%m%d_%H%M%S')}"
            score_dir = result / "01_누적성적표"
            wrong_dir = result / "02_회차별_개인오답"
            student_dir = result / "03_학생별_묶음"
            score_dir.mkdir(parents=True); wrong_dir.mkdir(); student_dir.mkdir()
            self.last_result = result

            self._progress(5)
            self._run_score_engine(selected, score_dir)
            self._progress(45)
            self._create_wrong_notes(selected, wrong_dir, student_dir)
            self._copy_score_images_to_students(score_dir, student_dir)
            self._progress(100)
            self._log(f"전체 완료: {result}")
            self.after(0, lambda: messagebox.showinfo("완료", f"성적표와 개인별 오답노트 생성이 완료되었습니다.\n\n{result}"))
        except Exception as e:
            self._log("오류: " + str(e))
            self.after(0, lambda: messagebox.showerror("오류", str(e)))
        finally:
            self.running = False
            self.after(0, lambda: self.btn.configure(state="normal"))

    def _open_result(self):
        p = self.last_result or Path(self.output_dir.get())
        p.mkdir(parents=True, exist_ok=True)
        if os.name == "nt": os.startfile(str(p))


def run_score_engine(argv: list[str]) -> int:
    if len(argv) != 3:
        return 2
    in_dir, out_dir, count = argv
    b = bundle_dir()
    score_dir = b / "score"
    if not score_dir.exists():
        score_dir = Path(__file__).resolve().parent / "score"
    old_argv = sys.argv[:]
    old_path = list(sys.path)
    try:
        sys.path.insert(0, str(score_dir))
        for script in ("run3.py", "build_xlsx.py"):
            target = score_dir / script
            sys.argv = [str(target), in_dir, out_dir, count]
            runpy.run_path(str(target), run_name="__main__")
        return 0
    finally:
        sys.argv = old_argv
        sys.path[:] = old_path


def main():
    if len(sys.argv) >= 2 and sys.argv[1] == "--score-engine":
        raise SystemExit(run_score_engine(sys.argv[2:]))
    App().mainloop()


if __name__ == "__main__":
    main()
