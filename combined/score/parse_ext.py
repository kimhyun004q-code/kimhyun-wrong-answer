# -*- coding: utf-8 -*-
"""기존 parse.py를 감싸서 숙제/전체공지/개인공지 데이터를 추가한다."""
import importlib.util
import os
import sys
import openpyxl

_HERE = os.path.dirname(os.path.abspath(__file__))
_BASE_PATH = os.path.join(_HERE, "parse_base.py")
_spec = importlib.util.spec_from_file_location("_parse_base", _BASE_PATH)
_base = importlib.util.module_from_spec(_spec)
sys.modules["_parse_base"] = _base
_spec.loader.exec_module(_base)

AREA_DESC = _base.AREA_DESC
AREA_ORDER = _base.AREA_ORDER
TEACHER = getattr(_base, "TEACHER", "김현T")

_PERSONAL_HEADERS = {"개인상담", "개인공지", "개인 공지", "개별공지", "개별 공지", "전달사항", "메모"}
_GLOBAL_HEADERS = {"전체공지", "전체 공지", "공통공지", "공통 공지", "공지사항"}


def _clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in {"None", "nan"} else s


def _read_notice_data(xlsx_path):
    global_notice = ""
    personal = {}
    try:
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        if "명단" not in wb.sheetnames:
            wb.close(); return global_notice, personal
        ws = wb["명단"]

        # 학생 헤더 행 찾기
        header_row = None
        for r in range(1, min(ws.max_row, 40) + 1):
            vals = [_clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 25) + 1)]
            if "학생명" in vals:
                header_row = r
                break
        if header_row is None:
            header_row = 10

        # 전체공지: 라벨 오른쪽 또는 라벨 아래쪽 영역을 읽는다.
        for r in range(1, min(header_row, 20) + 1):
            for c in range(1, min(ws.max_column, 25) + 1):
                raw = _clean(ws.cell(r, c).value)
                if raw in _GLOBAL_HEADERS or any(raw.startswith(h + "\n") for h in _GLOBAL_HEADERS):
                    parts = []
                    if "\n" in raw:
                        tail = raw.split("\n", 1)[1].strip()
                        if tail: parts.append(tail)
                    for cc in range(c + 1, min(ws.max_column, c + 4) + 1):
                        v = _clean(ws.cell(r, cc).value)
                        if v: parts.append(v)
                    for rr in range(r + 1, header_row):
                        v = _clean(ws.cell(rr, c).value)
                        if v: parts.append(v)
                    # 중복 제거, 입력 순서 유지
                    seen = set(); clean_parts = []
                    for x in parts:
                        if x not in seen:
                            seen.add(x); clean_parts.append(x)
                    global_notice = "\n".join(clean_parts).strip()
                    break
            if global_notice:
                break

        # 학생별 개인공지/개인상담 열 찾기
        name_col = None
        notice_col = None
        for c in range(1, min(ws.max_column, 30) + 1):
            h = _clean(ws.cell(header_row, c).value)
            if h == "학생명":
                name_col = c
            if h in _PERSONAL_HEADERS:
                notice_col = c
        if name_col and notice_col:
            for r in range(header_row + 1, ws.max_row + 1):
                nm = _clean(ws.cell(r, name_col).value)
                if not nm:
                    continue
                nm = _base.normalize_student_name(nm)
                if nm in _base.EXCLUDE_NAMES:
                    continue
                msg = _clean(ws.cell(r, notice_col).value)
                if msg:
                    personal[nm] = msg
        wb.close()
    except Exception:
        pass
    return global_notice, personal


def parse_session(date_key, xlsx_path, hwp_path=None, fallback_label=None, name_label=None):
    sess = _base.parse_session(date_key, xlsx_path, hwp_path, fallback_label, name_label)
    global_notice, personal = _read_notice_data(xlsx_path)
    sess["전체공지"] = global_notice
    for nm, at in sess.get("출결", {}).items():
        at["개인공지"] = personal.get(nm, "")
    return sess


def build(sessions, min_sessions=2):
    out = _base.build(sessions, min_sessions=min_sessions)
    for st in out:
        hw_rows = []
        global_rows = []
        personal_rows = []
        for r in st.get("rows", []):
            if r.get("결시"):
                continue
            s = r["sess"]
            label = s.get("hwp", {}).get("차시", "")
            hw = _clean(r.get("숙제"))
            if hw:
                hw_rows.append((label, hw))
            gn = _clean(s.get("전체공지"))
            if gn:
                global_rows.append((label, gn))
            at = s.get("출결", {}).get(st["이름"], {})
            pn = _clean(at.get("개인공지"))
            if pn:
                personal_rows.append((label, pn))
        st["숙제내역"] = hw_rows
        st["전체공지내역"] = global_rows
        st["개인공지내역"] = personal_rows
    return out
