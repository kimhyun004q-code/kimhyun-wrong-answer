# -*- coding: utf-8 -*-
"""4회 통합 성적표 - 데이터 파싱/집계 모듈"""
import os, re
from collections import defaultdict, Counter
import openpyxl, olefile

EXCLUDE_NAMES = {"평균", "최고점", "최저점", "이름없음", "이름없음2"}
HW_CODE = {"◎": 4, "○": 3, "△": 2, "미제출": 1, "매우잘함": 4, "잘함": 3, "부족": 2, "안함": 1, "〓": None}
ABSENT_MARKS = {"결석", "결", "무단", "무단결석", "미출석", "불참", "미응시", "x", "×", "✕", "✗", "-", "–", "—", "/"}


def normalize_student_name(v):
    s = str(v).strip()
    m = re.match(r"^(.*[가-힣])([A-Za-z]+)$", s)
    if m:
        return m.group(1) + m.group(2).upper()
    return s


def is_absent(v):
    return v is not None and str(v).strip().lower() in ABSENT_MARKS


def to_score(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def hwp_meta(path):
    if not path or not os.path.exists(path):
        return {"차시": "", "시험지": ""}
    f = olefile.OleFileIO(path)
    txt = f.openstream("PrvText").read().decode("utf-16-le", "ignore")
    head = txt[:200]
    m = re.search(r"<(\d+)\s*차시>", head)
    cha = f"{m.group(1)}차시" if m else ""
    m2 = re.search(r"<([^<>]*?T[^<>]*?)>", head)
    title = m2.group(1).strip() if m2 else ""
    return {"차시": cha, "시험지": title}


def find_row(ws, col, value, limit=30):
    for r in range(1, min(ws.max_row, limit) + 1):
        if ws.cell(r, col).value == value:
            return r
    return None


def parse_session(date_key, xlsx_path, hwp_path=None, fallback_label=None, name_label=None):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["명단"]
    dr = find_row(ws, 2, "날짜")
    info = {"날짜": ws.cell(dr, 3).value, "학원": ws.cell(dr + 1, 3).value, "반": ws.cell(dr + 2, 3).value}
    hr = find_row(ws, 3, "연락처") or find_row(ws, 2, "학생명")
    attend = {}
    for r in range(hr + 1, ws.max_row + 1):
        nm = ws.cell(r, 2).value
        if not nm or str(nm).strip() in EXCLUDE_NAMES:
            continue
        nm = normalize_student_name(nm)
        attend[nm] = {"출석": ws.cell(r, 4).value, "숙제": ws.cell(r, 5).value}

    ws = wb["테스트정보"]
    hr = find_row(ws, 2, "문제번호")
    qinfo = {}
    for r in range(hr + 1, ws.max_row + 1):
        qno, qtype, qarea = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if qtype is None or str(qtype).strip() == "":
            continue
        qinfo[int(qno)] = {"유형": str(qtype).strip(), "영역": str(qarea).strip()}
    qnos = sorted(qinfo)
    nq = len(qnos)

    ws = wb["테스트"]
    hr = find_row(ws, 4, "이름")
    qcol = {}
    for c in range(6, ws.max_column + 1):
        v = ws.cell(hr, c).value
        if isinstance(v, int):
            qcol[v] = c

    students = {}
    noscore = {}
    for r in range(hr + 1, ws.max_row + 1):
        nm = ws.cell(r, 4).value
        if not nm:
            continue
        nm = normalize_student_name(nm)
        if nm in EXCLUDE_NAMES:
            continue
        rank = ws.cell(r, 2).value
        if isinstance(rank, str) and "#" in rank:
            continue
        raw_score = ws.cell(r, 5).value
        score = to_score(raw_score)
        if score is None:
            noscore[nm] = "" if raw_score is None else str(raw_score).strip()
            continue
        wrong = [q for q in qnos if ws.cell(r, qcol[q]).value == 1]
        students[nm] = {"점수": score, "오답": wrong, "원등수": rank}

    scores = [s["점수"] for s in students.values()]
    n = len(scores)
    if n == 0:
        raise ValueError(f"[중단] {os.path.basename(xlsx_path)} 에 유효한 시험 점수가 한 건도 없습니다.")
    cls = {"응시자수": n, "평균": sum(scores) / n, "최고": max(scores), "최저": min(scores)}
    srt = sorted(scores, reverse=True)
    for nm, s in students.items():
        s["등수"] = srt.index(s["점수"]) + 1
        s["상위%"] = (s["등수"] - 1) / (n - 1) * 100 if n > 1 else 0.0

    qrate = {}
    for q in qnos:
        w = sum(1 for s in students.values() if q in s["오답"])
        qrate[q] = (n - w) / n

    ar = find_row(ws, 1, "응시", limit=ws.max_row)
    sheet_avg = ws.cell(ar, 5).value if ar else None
    top_types = [t for t, _ in Counter(qinfo[q]["유형"] for q in qnos).most_common(2)]
    hm = hwp_meta(hwp_path)
    if name_label:
        hm["차시"] = name_label
    elif not hm["차시"]:
        hm["차시"] = fallback_label or f"{int(date_key[2:4])}/{int(date_key[4:6])}"
    return {"key": date_key, "info": info, "hwp": hm, "문항": qinfo, "문항번호": qnos, "문항수": nq,
            "배점": 100.0 / nq, "정답률": qrate, "학생": students, "반": cls, "출결": attend,
            "미응시": noscore, "시트평균": sheet_avg, "대표유형": top_types}


NO_TREND = "–"
AREA_ORDER = ["계산", "이해", "추론", "문제해결", "자료해석"]
AREA_DESC = {"계산": "정의를 이용한 단순계산", "이해": "수업내용의 이해를 바탕으로 한 고난도 문제에의 적용",
             "추론": "조건에 따른 논리 전개", "문제해결": "고난도 사고력이 요구되는 문제", "자료해석": "그래프, 표, 그림, 함수의 변화를 읽고 판단하는 문제"}
AREA_NOTE = "  ·  ".join(f"{a} = {AREA_DESC[a]}" for a in AREA_ORDER)


def classify(rate):
    if rate >= 0.80:
        return "실수"
    if rate >= 0.50:
        return "개인취약"
    return "공통오답문제"


def major_unit_name(qtype):
    t = re.sub(r"\s+", " ", str(qtype or "")).strip()
    if not t:
        return ""
    for area_name in AREA_ORDER:
        for suffix in (f"의 {area_name}", f" {area_name}"):
            if t.endswith(suffix) and len(t) > len(suffix):
                t = t[:-len(suffix)].strip()
                break
    if t.count("의") >= 2:
        head, sep, tail = t.rpartition("의")
        if head.strip() and tail.strip():
            t = head.strip()
    return t.rstrip("의 ").strip()


def absence_reason(s, nm):
    if is_absent(s.get("미응시", {}).get(nm)):
        return "결석"
    at = s["출결"].get(nm)
    if at is None:
        return "결석"
    if is_absent(at.get("출석")):
        return "결석"
    return "결시"


def build(sessions, min_sessions=2):
    names = set()
    for s in sessions:
        names |= set(s["학생"])
    out = []
    for nm in sorted(names):
        rows = []
        for s in sessions:
            if nm not in s["학생"]:
                rows.append({"결시": True, "사유": absence_reason(s, nm), "sess": s})
                continue
            st = s["학생"][nm]
            at = s["출결"].get(nm, {})
            rows.append({"결시": False, "사유": "", "sess": s, "점수": st["점수"], "등수": st["등수"],
                         "상위%": st["상위%"], "오답": st["오답"], "출석": at.get("출석"), "숙제": at.get("숙제")})
        taken = [r for r in rows if not r["결시"]]
        if len(taken) < min_sessions:
            continue
        sc = [r["점수"] for r in taken]
        avg = sum(sc) / len(sc)
        dev = [r["점수"] - r["sess"]["반"]["평균"] for r in taken]
        h = len(dev) // 2
        if h >= 1:
            d = sum(dev[-h:]) / h - sum(dev[:h]) / h
            trend = "상승" if d >= 7 else ("하락" if d <= -7 else "유지")
            raw_d = sum(sc[-h:]) / h - sum(sc[:h]) / h
        else:
            d = raw_d = 0.0
            trend = NO_TREND

        area = defaultdict(lambda: {"n": 0, "correct": 0, "cls": 0.0})
        weak_key = Counter()
        weak_detail = defaultdict(list)
        wrongs = []
        for r in taken:
            s = r["sess"]
            for q in s["문항번호"]:
                a = s["문항"][q]["영역"]
                area[a]["n"] += 1
                area[a]["cls"] += s["정답률"][q]
                if q not in r["오답"]:
                    area[a]["correct"] += 1
            for q in r["오답"]:
                t = s["문항"][q]["유형"]
                a = s["문항"][q]["영역"]
                major = major_unit_name(t)
                key = (major, a)
                weak_key[key] += 1
                weak_detail[key].append((s["hwp"]["차시"], q, s["정답률"][q]))
                wrongs.append({"차시": s["hwp"]["차시"], "번호": q, "유형": t, "영역": a,
                               "반정답률": s["정답률"][q], "분류": classify(s["정답률"][q])})
        areas = []
        for a in AREA_ORDER:
            if a in area and area[a]["n"]:
                v = area[a]
                areas.append({"영역": a, "n": v["n"], "본인": v["correct"] / v["n"] * 100, "반": v["cls"] / v["n"] * 100})
        weak = []
        for (major, a), c in weak_key.most_common():
            if c < 2:
                continue
            label = f"{major}의 {a}" if major else a
            weak.append({"유형": label, "횟수": c, "대단원": major, "평가영역": a, "상세": weak_detail[(major, a)]})
            if len(weak) >= 3:
                break

        def code_avg(k):
            vals = [HW_CODE.get(str(r[k]).strip()) for r in taken if r[k] is not None]
            vals = [v for v in vals if v is not None]
            return sum(vals) / len(vals) if vals else None

        out.append({"이름": nm, "rows": rows, "응시": len(taken), "평균": avg, "최고": max(sc), "최저": min(sc),
                    "추이": trend, "추이차": d, "점수변화": raw_d, "편차평균": sum(dev) / len(dev), "편차": dev,
                    "평균상위%": sum(r["상위%"] for r in taken) / len(taken),
                    "반평균": sum(r["sess"]["반"]["평균"] for r in taken) / len(taken), "영역": areas, "취약": weak,
                    "오답": wrongs, "출석평균": code_avg("출석"), "숙제평균": code_avg("숙제"),
                    "분류집계": Counter(w["분류"] for w in wrongs)})
    return out
