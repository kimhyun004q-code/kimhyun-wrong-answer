# -*- coding: utf-8 -*-
import os, shutil, zipfile, glob
from collections import Counter

def _load_local_module(name):
    import importlib.util, sys
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), name + ".py")
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod

parse = _load_local_module("parse4")
report3 = _load_local_module("report4")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from playwright.sync_api import sync_playwright
from PIL import Image
import sys, discover

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

UP = sys.argv[1]
OUT = sys.argv[2]
MAX_SESSIONS = int(sys.argv[3]) if len(sys.argv) > 3 else 4
os.makedirs(OUT, exist_ok=True)
SESSIONS = discover.discover(UP, MAX_SESSIONS)
if not SESSIONS:
    sys.exit(f"[중단] {UP} 에서 회차 파일을 찾지 못했습니다.")
_m = discover.read_meta(SESSIONS[0][1])
S = [parse.parse_session(k, x, h, f"{i}회차", lb) for i, (k, x, h, lb) in enumerate(SESSIONS, 1)]
MIN_SESSIONS = 1
students = parse.build(S, min_sessions=MIN_SESSIONS)

import re as _re
_t = ""
for s_ in S:
    m = _re.search(r"([가-힣]+\s*T)", s_["hwp"]["시험지"] or "")
    if m and not _t:
        _t = m.group(1).replace(" ", "")
_d = lambda t: _re.sub(r"(\d{4})년 (\d+)월 (\d+)일.*", lambda g: f"{g.group(1)}.{int(g.group(2)):02d}.{int(g.group(3)):02d}", str(t))
CHA = [x["hwp"]["차시"] for x in S]
META = {"학원": _m.get("학원", ""), "반명": _m.get("반", ""), "강사": _t or report3.TEACHER,
        "기간": f'{_d(S[0]["info"]["날짜"])} – {_d(S[-1]["info"]["날짜"])[5:]}',
        "범위": f'{CHA[0]}–{CHA[-1]}' if len(S) > 1 else CHA[0]}
BASE = f'성적표 {META["반명"] or "반"} {SESSIONS[0][0]}-{SESSIONS[-1][0]}'

import openpyxl as _ox
_wbp = f"{OUT}/{BASE}.xlsx"
if os.path.exists(_wbp):
    try:
        _w = _ox.load_workbook(_wbp, data_only=True)
        _orig = report3.make_comment
        _cache = {}
        import hashlib
        _h = lambda L: hashlib.md5("\n".join(L).encode("utf-8")).hexdigest()[:12]
        for st in students:
            if st["이름"] in _w.sheetnames:
                _ws = _w[st["이름"]]
                ls = [_ws.cell(27 + i, 1).value for i in range(4)]
                ls = [l for l in ls if isinstance(l, str) and l.strip()]
                if ls and _h(ls) != _ws.cell(27, 40).value:
                    _cache[st["이름"]] = ls
        if _cache:
            report3.make_comment = lambda st: _cache.get(st["이름"], _orig(st))
            print(f"고쳐 쓰신 지도 방향 {len(_cache)}명분을 이미지에 반영합니다.")
    except Exception as e:
        print("  (워크북 코멘트 읽기 건너뜀:", e, ")")

print("=== 검증 ===")
for s in S:
    chk = f"시트 {s['시트평균']:.2f}" if isinstance(s['시트평균'], (int, float)) else "시트값 없음"
    print(f"[{s['key']} {s['hwp']['차시']}] 응시 {s['반']['응시자수']}명 평균 {s['반']['평균']:.2f} ({chk}) 문항 {s['문항수']}개")
_all = set().union(*[set(x["학생"]) for x in S])
print(f"전체 실명 {len(_all)}명 → {MIN_SESSIONS}회 이상 응시 {len(students)}명")

hp = f"{OUT}/{BASE}.html"
open(hp, "w", encoding="utf-8").write(report3.build_html(students, META))
imgdir = f"{OUT}/{BASE} 이미지"
shutil.rmtree(imgdir, ignore_errors=True)
os.makedirs(imgdir, exist_ok=True)
for _old in glob.glob(os.path.join(imgdir, "*.png")):
    try: os.remove(_old)
    except OSError: pass

made = []
with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": report3.CARD_W, "height": 900}, device_scale_factor=report3.SCALE)
    pg.goto("file://" + hp, wait_until="domcontentloaded")
    pg.wait_for_timeout(1000)
    cards = pg.query_selector_all(".card")
    assert len(cards) == len(students), f"카드 {len(cards)} != 학생 {len(students)}"
    seen = {}
    for i, (el, st) in enumerate(zip(cards, students), 1):
        shown = el.query_selector("h1").inner_text().strip()
        assert shown == st["이름"], f"{i}번째 카드 이름 불일치: {shown} != {st['이름']}"
        seen[shown] = seen.get(shown, 0) + 1
        sfx = f" ({seen[shown]})" if seen[shown] > 1 else ""
        path = f"{imgdir}/{i:02d} {shown}{sfx}.png"
        el.screenshot(path=path)
        made.append(path)
    b.close()
if not made:
    raise RuntimeError("성적표 이미지가 생성되지 않았습니다.")

sizes = [Image.open(f).size for f in made]
print(f"이미지 {len(made)}장 · 폭 {sizes[0][0]}px")
with zipfile.ZipFile(f"{OUT}/{BASE} 이미지.zip", "w", zipfile.ZIP_DEFLATED) as z:
    for f in made: z.write(f, os.path.basename(f))
ims = [Image.open(f).convert("RGB") for f in made]
ims[0].save(f"{OUT}/{BASE}.pdf", save_all=True, append_images=ims[1:], resolution=72 * report3.SCALE)

wb = openpyxl.Workbook()
ws = wb.active; ws.title = "점수매트릭스"
ws.append(["이름", "응시"] + [f"{s['hwp']['차시']}({s['key']})" for s in S] + ["평균", "최고", "최저", "평균상위%", "반대비편차", "추이"])
for st in students:
    ws.append([st["이름"], st["응시"]] + [(r.get("사유") or "결시") if r["결시"] else r["점수"] for r in st["rows"]] +
              [round(st["평균"], 1), st["최고"], st["최저"], round(st["평균상위%"], 1), round(st["편차평균"], 1), st["추이"]])
ws.append(["반 평균", ""] + [round(s["반"]["평균"], 1) for s in S])
ws2 = wb.create_sheet("유형별오답")
ws2.append(["회차", "문항", "유형", "영역", "반 정답률", "오답자수", "오답자"])
for s in S:
    for q in s["문항번호"]:
        w = [n for n, v in s["학생"].items() if q in v["오답"]]
        ws2.append([s["hwp"]["차시"], q, s["문항"][q]["유형"], s["문항"][q]["영역"], round(s["정답률"][q], 3), len(w), ", ".join(sorted(w))])
ws3 = wb.create_sheet("학생별취약유형")
ws3.append(["이름", "반복취약유형", "실수", "개인취약", "공통오답문제", "총오답"])
for st in students:
    c = st["분류집계"]
    ws3.append([st["이름"], " / ".join(w["유형"] for w in st["취약"]) or "-", c.get("실수", 0), c.get("개인취약", 0), c.get("공통오답문제", 0), len(st["오답"])])
hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="10203A")
for w in (ws, ws2, ws3):
    for cell in w[1]:
        cell.font = hf; cell.fill = hfill; cell.alignment = Alignment(horizontal="center")
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 9), 45)
    w.freeze_panes = "A2"
wb.save(f"{OUT}/{BASE} 집계.xlsx")
print("회차별 반 평균: " + " → ".join(f"{s['hwp']['차시']} {s['반']['평균']:.1f}" for s in S))
