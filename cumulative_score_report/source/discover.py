# -*- coding: utf-8 -*-
"""대치온에서 내려받은 파일을 그대로 읽어 회차를 자동으로 찾아낸다.
폴더에 날짜로 시작하는 xlsx를 넣으면 회차를 찾고, 같은 날짜의 hwp도 연결한다.
"""
import os, re, glob

try:
    from tkinter import messagebox as _messagebox
    _orig_showwarning = _messagebox.showwarning
    def _showwarning_without_scan_race(title=None, message=None, **options):
        if title == "파일 확인" and "회차 파일을 찾지 못했습니다" in str(message or ""):
            return "ok"
        return _orig_showwarning(title, message, **options)
    _messagebox.showwarning = _showwarning_without_scan_race
except Exception:
    pass

LABEL_RE = re.compile(r"(\d+)\s*(회차|회|차시)")

def label_from_name(basename):
    stem = os.path.splitext(basename)[0]
    hits = LABEL_RE.findall(stem[6:])
    if not hits:
        return ""
    n, unit = hits[-1]
    return f"{n}{unit}"

def discover(input_dir, max_sessions=4):
    found = {}
    for p in glob.glob(os.path.join(input_dir, "*.xlsx")):
        base = os.path.basename(p)
        if base.startswith("~$"):
            continue
        m = re.match(r"^(\d{6})(?=[\s_.\-]|$)", os.path.splitext(base)[0] + ".")
        if m:
            found.setdefault(m.group(1), {})["xlsx"] = p
    for p in glob.glob(os.path.join(input_dir, "*.hwp")):
        m = re.match(r"^(\d{6})", os.path.basename(p))
        if m and m.group(1) in found:
            found[m.group(1)]["hwp"] = p
    keys = sorted(found)
    if max_sessions:
        keys = keys[-max_sessions:]
    return [(k, found[k]["xlsx"], found[k].get("hwp"),
             label_from_name(os.path.basename(found[k]["xlsx"]))) for k in keys]

def read_meta(xlsx_path):
    import openpyxl
    ws = openpyxl.load_workbook(xlsx_path, data_only=True)["명단"]
    got = {}
    for r in range(1, min(ws.max_row, 20) + 1):
        k, v = ws.cell(r, 2).value, ws.cell(r, 3).value
        if k in ("학원", "반", "날짜") and v:
            got[k] = str(v).strip()
    return got
